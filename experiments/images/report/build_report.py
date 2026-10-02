#!/usr/bin/env python3
"""Build the German Clef vision report from a completed, frozen benchmark.

This file is deliberately outside the frozen benchmark. It never imports or
loads a model, starts inference, alters scoring, or fabricates missing results.
Render the resulting DOCX to PDF, then inspect every
page before delivery. Two expressly authorized documentary totals crops are
embedded losslessly; no complete source image is embedded.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
FIELDS = [
    ("chart", "chart_type", "Diagrammart"),
    ("chart", "legend_count", "Legendenanzahl"),
    ("invoice", "document_type", "Dokumentart"),
    ("invoice", "tax_note", "Steuerhinweis"),
    ("invoice", "gross_band", "Bruttosummenintervall"),
]
LABELS = {
    "bar_line": "Säulen mit Linie", "bar_pie": "Kreis mit gestapelter Säule",
    "hbar": "gruppierte horizontale Balken", "hbar2": "horizontale Balken mit Vorzeichenlegende",
    "line": "Linien", "pie": "Kreis", "stack_hbar": "gestapelte horizontale Balken",
    "stack_vbar": "gestapelte Säulen", "vbar": "gruppierte Säulen mit einer y-Achse",
    "vbar2": "Säulen mit zwei y-Achsen", "rechnung": "Rechnung",
    "gutschrift": "Gutschrift", "regular": "regulärer Umsatzsteuerhinweis",
    "small_business": "Kleinunternehmerhinweis", "reverse_charge": "Reverse Charge",
    "negative": "negativ", "zero_1000": "0 bis einschließlich 1.000 EUR",
    "1000_5000": "über 1.000 bis einschließlich 5.000 EUR",
    "5000_20000": "über 5.000 bis einschließlich 20.000 EUR", "over_20000": "über 20.000 EUR",
}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def read_jsonl(path):
    return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def percent(value, decimals=1):
    require(isinstance(value, (int, float)), "Missing numeric percentage")
    return f"{100 * value:.{decimals}f}".replace(".", ",") + " %"


def pp(value):
    return f"{100 * value:+.1f}".replace(".", ",") + " PP"


def decimal(value, digits=1):
    return f"{value:.{digits}f}".replace(".", ",")


def result_text(metric):
    return f"{metric['correct']}/{metric['total']} ({percent(metric['accuracy'])})"


def load_checked(root, scores_path, metadata_path):
    import sys
    sys.path.insert(0,str(root/'scripts'))
    from provenance_checks import matches
    """Refuse partial, mixed-version or modified benchmark inputs."""
    missing = [str(p) for p in (scores_path, metadata_path) if not p.is_file()]
    if missing:
        raise FileNotFoundError("Final result input not yet available: " + ", ".join(missing))
    s, m = read_json(scores_path), read_json(metadata_path)
    require(m.get("status") == "completed", "Inference metadata is not completed")
    require(m.get("request_count") == 90, "Expected exactly 90 scored forward runs")
    require(m.get("max_pixels") == 786432, "Unexpected pixel limit")
    require(m.get("vision_linear_quantized_count") == 0, "Vision quantization differs from specification")
    require(m.get("vision_parameter_dtypes") == ["torch.bfloat16"], "Vision dtype differs from specification")
    require(m.get("joint_head_dtypes") == ["torch.bfloat16"], "Joint-head dtype differs from specification")
    require(m.get("output_embedding_dtype") == "torch.bfloat16", "Output dtype differs from specification")
    require(m.get("no_gold_loaded") is True, "No-gold input assurance missing")
    require(m.get("batch_size") == 1 and m.get("threads") == 6, "Unexpected execution configuration")
    require(m["quantization"].get("bnb_4bit_quant_type") == "nf4", "NF4 configuration missing")
    require(m["quantization"].get("bnb_4bit_use_double_quant") is True, "Double quantization missing")
    freeze_path = root / "benchmark/freeze_manifest.json"
    freeze = read_json(freeze_path)
    for relative, expected in freeze["files_sha256"].items():
        require(matches(root,relative,expected), "Frozen input changed: " + relative)
    require(sha(freeze_path) == s["provenance"]["freeze_manifest_sha256"], "Scores use another freeze")
    require(matches(root,"scripts/score_images.py",s["provenance"]["scorer_sha256"]), "Scorer hash differs")
    require(sha(root / "results/predictions.jsonl") == s["provenance"]["predictions_sha256"], "Predictions hash differs")
    require(sha(root / "benchmark/requests.jsonl") == m["requests_sha256"], "Request hash differs")
    require(matches(root,"scripts/run_images.py",m["runner_sha256"]), "Runner hash differs")
    require(s["schema_total"] == 90 and len(s["records"]) == 90, "Incomplete score records")
    require(len({r["id"] for r in s["records"]}) == 90, "Duplicate scored records")
    require(set(m["request_order_ids"]) == {r["id"] for r in s["records"]}, "Scored request set differs")
    main = [f for f in s["fields"] if f["language"] == "de" and f["condition"] == "image"]
    require(len(main) == 120 and len({f["case_id"] for f in main}) == 50, "Primary field/image count differs")
    require(sum(x["paired_images"] for x in s["paired_controls"].values()) == 20, "Paired image count differs")
    for kind, field, _ in FIELDS:
        group = s["groups"][kind + "_de_image"]
        expected_n = 30 if kind == "chart" else 20
        metric = group["per_field"][field]
        require(metric["total"] == expected_n, "Unexpected primary task denominator: " + field)
        rows = [f for f in main if f["field"] == field]
        require(metric["correct"] == sum(f["correct"] for f in rows), "Task total inconsistent: " + field)
        require(abs(metric["accuracy"] - metric["correct"] / expected_n) < 1e-10, "Task ratio inconsistent")
        pairs = [r for r in s["paired_controls"][kind]["rows"] if r["field"] == field]
        require(len(pairs) == 10, "Expected ten paired observations per task")
    design = read_json(root / "benchmark/design_summary.json")
    require((design["images"], design["requests"], design["main_fields"]) == (50, 90, 120), "Design changed")
    return s, m, design, read_json(root / "source/provenance.json")


def xml(parent, name, **attrs):
    el = OxmlElement("w:" + name)
    for key, val in attrs.items():
        el.set(qn("w:" + key), str(val))
    parent.append(el)
    return el


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.68)
    section.left_margin = section.right_margin = Inches(.72)
    section.footer_distance = Inches(.3)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(26)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.styles["Subtitle"].font.size = Pt(11)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(15)
    for name, size in [("Heading 1", 17), ("Heading 2", 11.5)]:
        st = doc.styles[name]
        st.font.size, st.font.bold = Pt(size), True
        st.paragraph_format.space_before = Pt(13 if name == "Heading 2" else 0)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = False
    settings = doc.settings.element
    xml(settings, "updateFields", val="true")
    normal = doc.styles["Normal"].element.get_or_add_rPr()
    xml(normal, "lang", val="de-DE")
    for pstyle in [doc.styles["Title"], doc.styles["Subtitle"], doc.styles["Heading 1"], doc.styles["Heading 2"]]:
        ppr = pstyle.element.find(qn("w:pPr"))
        if ppr is not None:
            for border in ppr.findall(qn("w:pBdr")):
                ppr.remove(border)
    # A multi-page benchmark needs pagination for citations and review.
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = foot.add_run("Clef Bildbenchmark  |  ")
    run.font.name, run.font.size = "Arial", Pt(8)
    for typ, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        run = foot.add_run()
        if typ:
            xml(run._r, "fldChar", fldCharType=typ)
        else:
            instr = xml(run._r, "instrText")
            instr.text = text
    doc.core_properties.title = "Clef Bildbenchmark mit deutschen Fragen"
    doc.core_properties.subject = "Ergebnisse eines kontrollierten synthetischen Bildtests"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "Clef, Vision, Benchmark, Deutsch, synthetische Daten"


def para(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def table(doc, headers, rows, widths, font_size=9.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        for i, value in enumerate(row):
            if i == 0:
                cells = t.add_row().cells
            cells[i].text = str(value)
    props = t._tbl.tblPr
    borders = xml(props, "tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        xml(borders, side, val="single", sz=4, color="D9D9D9")
    margins = xml(props, "tblCellMar")
    for side, value in [("top", 75), ("bottom", 75), ("left", 90), ("right", 90)]:
        xml(margins, side, w=value, type="dxa")
    for row_idx, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        xml(trpr, "cantSplit")
        if row_idx == 0:
            xml(trpr, "tblHeader")
        for col_idx, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            xml(cell._tc.get_or_add_tcPr(), "shd", fill="243C50" if row_idx == 0 else ("F1F4F6" if row_idx % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.03
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = "Arial"
                    run.font.bold = row_idx == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_idx == 0 else "000000")
    para(doc, "").paragraph_format.space_after = Pt(0)
    return t


def link(doc, title, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    relationship = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), relationship)
    r = xml(h, "r")
    props = xml(r, "rPr")
    xml(props, "color", val="234E70")
    xml(props, "u", val="single")
    xml(props, "sz", val=20)
    xml(r, "t").text = title
    p._p.append(h)
    return p


def small(doc, text, mono=False):
    p = para(doc, text)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        r.font.size = Pt(8.5 if mono else 9)
        if mono:
            r.font.name = "Courier New"
    return p


def new_page(doc, title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def make_evidence(root, output_dir):
    """Lossless exact-pixel crops; source boxes anchor the totals context."""
    cases = {c["case_id"]: c for c in read_jsonl(root / "benchmark/cases.jsonl")}
    # Label boxes enclose values only. Fixed context extensions include the
    # adjacent Netto/USt/Gesamtbetrag labels and their original rules/box.
    extensions = {"invoice-009": (212, 26, 24, 23), "invoice-016": (466, 29, 19, 29)}
    records = []
    output_dir.mkdir(parents=True, exist_ok=True)
    for case_id, (left, top, right, bottom) in extensions.items():
        case = cases[case_id]
        source_path = root / case["image_path"]
        label_path = root / case["source_label_path"]
        require(sha(source_path) == case["image_sha256"], "Evidence image changed")
        require(sha(label_path) == case["source_label_sha256"], "Evidence label changed")
        label = read_json(label_path)
        keys = ["net_total", "vat_amount_19", "gross_total"]
        boxes = [label["fields"][key]["box"] for key in keys]
        union = [min(b[0] for b in boxes), min(b[1] for b in boxes),
                 max(b[2] for b in boxes), max(b[3] for b in boxes)]
        rect = [union[0]-left, union[1]-top, union[2]+right, union[3]+bottom]
        with Image.open(source_path) as source:
            require(source.size == tuple(label["image_size"]), "Evidence coordinate frame changed")
            cropped = source.crop(tuple(rect))
            destination = output_dir / f"{case_id}_totals.png"
            cropped.save(destination, format="PNG")
            with Image.open(destination) as saved:
                require(saved.size == cropped.size and saved.tobytes() == cropped.tobytes(),
                        "Evidence crop is not pixel-identical")
        records.append({
            "case_id": case_id, "source_id": case["source_id"],
            "source_revision": case["source_revision"],
            "source_image_sha256": case["image_sha256"],
            "source_label_sha256": case["source_label_sha256"],
            "crop_rect_xyxy": rect, "original_label_boxes": dict(zip(keys, boxes)),
            "crop_sha256": sha(destination), "crop_file": str(destination.relative_to(root)),
            "pixel_identical_to_source_region": True, "retouched": False,
            "amount_texts": {key: label["fields"][key]["text"] for key in keys},
            "attribution": "Belege (Immineal, 2026)",
        })
    return records


def create_report(s, m, design, provenance, root, destination, scores_path, metadata_path):
    doc = Document()
    configure(doc)
    groups = s["groups"]
    main = [f for f in s["fields"] if f["language"] == "de" and f["condition"] == "image"]
    errors = s["german_image_errors"]
    high_errors = s["german_image_errors_confidence_ge_0_9"]
    chart = groups["chart_de_image"]
    invoice = groups["invoice_de_image"]
    above = [label for kind, field, label in FIELDS if groups[kind + "_de_image"]["per_field"][field]["accuracy"] > groups[kind + "_de_image"]["majority_class_baseline_per_field"][field]]

    # Page 1: conclusion and all five primary metrics.
    doc.add_heading("Clef Bildbenchmark\nmit deutschen Fragen", 0)
    para(doc, "Ergebnisse auf synthetischen Diagrammen und deutschen Belegen\n2. Oktober 2026", "Subtitle")
    doc.add_heading("Fazit", 2)
    para(doc, f"Clef beantwortet bei {result_text(chart['all_fields_per_image'])} der Diagramme und "
         f"{result_text(invoice['all_fields_per_image'])} der Belege sämtliche Aufgaben richtig. "
         f"In {len(above)} der fünf Aufgaben liegt die Trefferquote über der jeweiligen Mehrheitsklassenbaseline. "
         "Die Ergebnisse beschreiben einen kleinen, kontrollierten Bildtest und reichen für eine Freigabe "
         "zur automatischen Dokumentenverarbeitung oder für allgemeine Finanzanwendungen nicht aus.")
    para(doc, "Umfang: 50 eigenständige Quellbilder, 120 Primärentscheidungen und 90 gewertete Vorwärtsläufe. "
         "Die 30 Diagramme sind synthetisch und englisch beschriftet. Die 20 Rechnungen und Gutschriften "
         "sind synthetische Originalbilder mit tatsächlich deutschem Text. Alle Hauptfragen sind deutsch.", bold_lead="Umfang:")
    para(doc, "Wichtige Einschränkung: Die drei strikten Diagrammartfehler betreffen potenziell überlappende "
         "Antwortoptionen. Sie belegen daher nicht eindeutig eine visuelle Fehlwahrnehmung. Der eingefrorene "
         "Score bleibt unverändert; die nachträglich erkannte Mehrdeutigkeit wird auf Seite 4 erläutert.",
         bold_lead="Wichtige Einschränkung:")
    doc.add_heading("Die fünf Hauptaufgaben", 2)
    primary_rows = []
    for kind, field, label in FIELDS:
        g = groups[kind + "_de_image"]
        a = g["per_field"][field]
        primary_rows.append([label, f"{a['correct']}/{a['total']}", percent(a["accuracy"]),
                             percent(g["balanced_accuracy_per_field"][field]),
                             percent(g["majority_class_baseline_per_field"][field])])
    table(doc, ["Aufgabe", "Richtig", "Trefferquote", "Balanced\nAccuracy", "Mehrheits-\nbaseline"],
          primary_rows, [2.26, .65, 1.13, 1.28, 1.18])
    small(doc, "Balanced Accuracy = mittlere Trefferquote über die im Test vorhandenen Referenzklassen. "
          "Die Mehrheitsbaseline wählt für jede Aufgabe stets deren häufigste Referenzklasse. "
          "Die beiden Maße verhindern, dass ungleiche Klassenhäufigkeiten unbemerkt bleiben.")
    doc.add_heading("Technische Gültigkeit und Unsicherheit", 2)
    high_total = sum(isinstance(f["confidence"], (int, float)) and f["confidence"] >= .9 for f in main)
    para(doc, f"Schema gültig: {s['schema_valid']}/{s['schema_total']} Ausgaben. "
         f"Tatsächlicher Vision-Aufruf bei jedem Request: {'ja' if s['all_expected_images_reached_vision_encoder'] else 'nein'}. "
         f"Gekürzte Eingaben: {s['truncated_records']}. "
         f"Unter den {high_total} Primärentscheidungen mit mindestens 90 % Modellkonfidenz liegen "
         f"{len(high_errors)} Fehler. Konfidenz ist eine Softmax-Ausgabe und keine nachgewiesene Verlässlichkeit.")

    # Page 2: fair paired comparison, never compare unmatched primary subsets.
    new_page(doc, "Gepaarter Sprachvergleich und Weißbildkontrolle")
    para(doc, "Für dieselben 20 ausgewählten Bilder werden drei Bedingungen verglichen: deutsches "
         "Fragenset mit Originalbild, englisches Fragenset mit Originalbild und deutsches Fragenset "
         "mit einer weißen Bildfläche in identischer Originalgröße. Je zehn Diagramme und zehn Belege "
         "liefern 50 Entscheidungen pro Bedingung. Die zusätzlichen 40 Läufe sind keine neuen Bilder.")
    paired_rows = []
    for kind, field, label in FIELDS:
        rows = [r for r in s["paired_controls"][kind]["rows"] if r["field"] == field]
        n = len(rows)
        de, en, blank = (sum(r[key] for r in rows) for key in ("de_correct", "en_correct", "blank_correct"))
        paired_rows.append([label, f"{de}/{n}", f"{en}/{n}", f"{blank}/{n}", pp((en-de)/n), pp((de-blank)/n)])
    table(doc, ["Aufgabe", "DE\nBild", "EN\nBild", "DE\nWeiß", "EN minus\nDE", "Bild minus\nWeiß"],
          paired_rows, [2.06, .72, .72, .72, 1.18, 1.10])
    small(doc, "Jede Zelle mit Treffern hat denselben Nenner 10. PP = Prozentpunkte. "
          "EN minus DE misst den Unterschied auf identischen Bildern; Bild minus Weiß hält die deutsche Frage konstant.")
    for kind, label in [("chart", "Diagramme"), ("invoice", "Deutsche Belege")]:
        p = s["paired_controls"][kind]
        rows, n = p["rows"], p["paired_fields"]
        en_only = sum(r["en_correct"] and not r["de_correct"] for r in rows)
        de_only = sum(r["de_correct"] and not r["en_correct"] for r in rows)
        doc.add_heading(label, 2)
        para(doc, f"Deutsch mit Bild: {p['de_correct']}/{n} ({percent(p['de_correct']/n)}); "
             f"Englisch mit Bild: {p['en_correct']}/{n} ({percent(p['en_correct']/n)}); "
             f"Deutsch mit Weißbild: {p['blank_correct']}/{n} ({percent(p['blank_correct']/n)}). "
             f"Nur die englische Variante ist in {en_only} Entscheidungen richtig, nur die deutsche in {de_only}. "
             f"Beide Sprachen wählen in {p['de_en_same_choices']}/{n} Entscheidungen dieselbe Option.")
        para(doc, f"Durch das Weißbild ändern sich {p['de_blank_changed_choices']}/{n} gewählte Optionen. "
             f"Die mittlere Wahrscheinlichkeit der Referenzantwort ändert sich von Weiß zu Bild um "
             f"{pp(p['mean_gold_probability_image_minus_blank'])}.")
    doc.add_heading("Was die Kontrollen aussagen", 2)
    para(doc, "Ein Vorsprung mit Originalbild ist ein Hinweis auf nutzbaren Bildinhalt. Gleichbleibende oder "
         "richtige Weißbildantworten können aus Klassen-, Options- und Sprachpräferenzen entstehen. "
         "Da keine Enthaltungsoption angeboten wird, erzwingt das Schema auch beim Weißbild eine Auswahl; "
         "solche Antworten werden nicht als Halluzinationen gewertet. Das Experiment trennt diese Mechanismen "
         "nicht vollständig und prüft weder andere Bildstörungen "
         "noch reale Robustheit. Sprachunterschiede bleiben wegen der kleinen, gezielt ausgewählten "
         "Teilmenge deskriptiv; der offizielle Rahmenprompt bleibt in beiden Sprachen englisch.")

    # Page 3: sampling, gold and execution conditions.
    new_page(doc, "Stichprobe und Auswertungsmethode")
    doc.add_heading("Deterministische Auswahl vor der Inferenz", 2)
    para(doc, "Diagramme: 30 von 300 Bildern des offiziellen Testsplit von YuukiAsuna/synthetic_chart. "
         "Für jede Kombination aus zehn Diagrammtypen und drei Schwierigkeitsgraden wurde das Bild mit "
         "dem kleinsten SHA-256-Wert gewählt. Die Felder sind Diagrammart und Anzahl sichtbarer Legendeneinträge.", bold_lead="Diagramme:")
    para(doc, "Belege: 20 aus der öffentlichen 40-Dokumente-Vorschau ohne offiziellen Testsplit. "
         "Die Auswahl umfasst acht reguläre, fünf Kleinunternehmer-, drei Reverse-Charge- und vier "
         "Gutschrift-Fälle; zwölf Fotos, sechs Scans und zwei saubere Renderings. Nicht reguläre Fälle "
         "sind vollständig enthalten, reguläre wurden nach Variante und Layout geschichtet und anhand "
         "von SHA-256 ausgewählt.", bold_lead="Belege:")
    para(doc, "Die drei Belegaufgaben prüfen Dokumentart, ausdrücklich gedruckten Steuerhinweis und das "
         "Intervall der Bruttosumme. Die fünf Intervalle sind: negativ; 0 bis einschließlich 1.000; "
         "über 1.000 bis einschließlich 5.000; über 5.000 bis einschließlich 20.000; über 20.000 EUR. "
         "Jede Intervallklasse kommt viermal vor.")
    doc.add_heading("Referenzen und Bewertung", 2)
    para(doc, "Alle 50 Bilder und 120 Referenzentscheidungen wurden vor dem Hauptlauf direkt an den "
         "Bildpunkten unabhängig geprüft, ohne Modellantworten zu lesen. Zwei Optionsbeschreibungen "
         "wurden vor dem Einfrieren präzisiert. Bilder, Fälle, Referenzen, Fragen und Scorer sind "
         "mit Hashwerten fixiert. Gewertet wird die exakte zulässige Options-ID; ungültige Felder "
         "gelten als falsch. Vollständig richtige Bilder erfordern richtige Antworten auf alle ihre Felder.")
    doc.add_heading("Ausführung auf CPU", 2)
    para(doc, f"Cloudflare/clef-flash, Revision {m['revision']}. Batch {m['batch_size']}, "
         f"{m['threads']} Threads. Sprach-Linear-Layer: NF4 mit Double Quantization und BF16-Berechnung. "
         "Originaler Vision-Encoder, Joint Head und Ausgabe-Embedding: BF16. Der offizielle "
         "Modellwrapper ist unverändert. Diese experimentelle CPU-Konfiguration ist nicht mit "
         "Herstellerangaben für H200 oder eine vollständige BF16-Ausführung gleichzusetzen.")
    para(doc, "Der offizielle Processor erhält 65.536 bis 786.432 Pixel unter Erhalt des Seitenverhältnisses "
         "und rundet auf sein Patchraster. Keine Ausschnitte, Rotation, Bildoptimierung oder zusätzliche OCR. "
         "Dateinamen, Quell-IDs, Labels, Metadaten und OCR-Transkripte gelangen nicht in den Modelltext. "
         "Ein Hook protokolliert den tatsächlichen Encoderaufruf; echte Bildtensoren werden erzeugt.")
    small(doc, "Der Vision-Encoder bleibt auch aus technischem Grund BF16: Die installierte CPU-4-Bit-Packroutine "
          "passt nicht zu seiner MLP-Dimension 4304. Vor dem Lauf wurden null quantisierte Vision-Linear-Module "
          "und ausschließlich BF16-Visionparameter geprüft. Die vollständige Kodierung wurde für alle 90 "
          "Requests vorab kontrolliert, ohne Kürzung.")

    # Page 4: observed errors and source defects are clearly distinguished.
    new_page(doc, "Fehler und Grenzen der Aussagekraft")
    doc.add_heading("Fehler mit hoher Modellkonfidenz", 2)
    error_rows = []
    for kind, field, label in FIELDS:
        ff = [f for f in main if f["field"] == field]
        ee = [f for f in ff if not f["correct"]]
        hh = [f for f in ee if isinstance(f["confidence"], (int, float)) and f["confidence"] >= .9]
        error_rows.append([label, f"{len(ee)}/{len(ff)}", len(hh)])
    table(doc, ["Aufgabe", "Fehler gesamt", "Fehler mit Konfidenz ≥ 90 %"], error_rows, [2.55, 1.3, 2.65])
    doc.add_heading("Mehrdeutige Diagrammartoptionen", 2)
    ce = [f for f in errors if f["field"] == "chart_type"]
    require(len(ce) == 3 and all(f["expected"] == "bar_line" and f["predicted"] == "vbar2" for f in ce),
            "The report's post-hoc ambiguity discussion needs review for these results")
    para(doc, "Alle drei Abweichungen (chart-001 bis -003) wählen „Senkrechte Säulen mit zwei y-Achsen“ "
         "statt „Kombination aus Säulen und Linie“. Die eingefrorene erste Beschreibung schließt Linien "
         "nicht aus und kann sich mit dem Kombinationsdiagramm überschneiden. Das ist eine potenzielle "
         "Benchmarkmehrdeutigkeit, kein bewiesener visueller Erkennungsfehler. Der bestandene "
         "Vorab-Pixelaudit erkannte sie nicht; sie fiel erst nach dem Lauf auf. Der strikte Wert 27/30 "
         "bleibt bestehen. Eingaben, Referenzen und Hauptlauf wurden nicht nachträglich geändert.")
    doc.add_heading("Vorzeichen und Steuerhinweis", 2)
    gross_errors = [f for f in errors if f["field"] == "gross_band"]
    require(len(gross_errors) == 2 and all(f["expected"] == "negative" for f in gross_errors),
            "The report's gross-error discussion needs review for these results")
    details = []
    for f in gross_errors:
        details.append(f"{f['case_id']}: „{LABELS[f['predicted']]}“ bei {percent(f['confidence'])}")
    para(doc, "Zwei negative Gutschrift-Gesamtbeträge werden positiven Intervallen zugeordnet: "
         + "; ".join(details) + ". Die Dokumentart Gutschrift wird in beiden Fällen richtig erkannt. "
         "Für negative Bruttobeträge sind nur zwei von vier Entscheidungen richtig.")
    tax_errors = [f for f in errors if f["field"] == "tax_note"]
    require(len(tax_errors) == 1, "The report's tax-error discussion needs review for these results")
    te = tax_errors[0]
    para(doc, f"{te['case_id']} wird trotz §19-Fußnote als regulär besteuert eingestuft "
         f"({percent(te['confidence'])} Konfidenz). Die Frage gibt dem ausdrücklichen Hinweis Vorrang "
         "vor den widersprüchlichen Steuersätzen in Positionszeilen. Unter den fünf "
         "Kleinunternehmerfällen sind vier Steuerhinweise richtig erkannt.")
    doc.add_heading("Quellfehler und Grenzen", 2)
    para(doc, "Vier Kleinunternehmerbelege (invoice-004, -006, -012 und -013) zeigen Umsatzsteuersätze in "
         "Positionszeilen trotz ausdrücklichem §19-Hinweis. Gewertet wird gemäß Frage dieser gedruckte "
         "Fußnotenhinweis, nicht steuerrechtliche Richtigkeit. Bei drei Kreisdiagrammen ist Legendentext "
         "abgeschnitten, aber jede zu zählende Farbbox und Zeile sichtbar. Synthetische Titel, Achsen "
         "und Kategorien sind teils inhaltlich unplausibel.")
    para(doc, "Die gezielte synthetische Stichprobe liefert keine Populationsschätzung; mögliche "
         "Trainingskontamination ist unbekannt. Mehrere Felder desselben Bildes und gepaarte Wiederholungen "
         "sind keine unabhängigen Beobachtungen. Verkleinerung kann Text unlesbarer machen. Nicht gemessen "
         "werden freie Volltext-OCR, komplexe Finanzmathematik, Rechnungsprüfung, Anlageberatung, "
         "Compliance oder sichere Verarbeitung realer vertraulicher Unterlagen.")

    # Page 5: source attribution, observed timings and executable provenance.
    new_page(doc, "Quellen und Reproduktion")
    sources = {src["repo"]: src for src in provenance["sources"]}
    chart_src = sources["YuukiAsuna/synthetic_chart"]
    invoice_src = sources["laterrr/belege-de-invoices-sample"]
    link(doc, "Modell Cloudflare clef flash", f"https://huggingface.co/{m['model']}/tree/{m['revision']}")
    small(doc, f"Revision {m['revision']} · Lizenz Apache 2.0")
    link(doc, "Diagramme YuukiAsuna synthetic chart", chart_src["source_url"])
    small(doc, f"Revision {chart_src['revision']} · offizieller Testsplit · CC BY 4.0 laut Anbieter")
    link(doc, "Deutsche Belege Belege von Immineal 2026", invoice_src["source_url"])
    small(doc, f"Revision {invoice_src['revision']} · kostenlose Vorschau ohne offiziellen Testsplit")
    para(doc, "Belege (Immineal, 2026). Die Lizenz erlaubt lokale Evaluation sowie die Veröffentlichung "
         "abgeleiteter Ergebnisse und dokumentarischer Abbildungen mit dieser Attribution. Dieser Bericht "
         "enthält zwei Summenausschnitte als Fehlerbelege, keine vollständigen Rechnungsbilder oder "
         "Weiterverbreitung des ausgewählten Datensatzes. Alle Personen, Firmen und Beträge der "
         "Belege sind laut Anbieter frei erfunden.")
    link(doc, "Belege Dataset Licence Version 1 vom 14 August 2026", invoice_src["license_url"])
    doc.add_heading("Laufzeit unter den gemessenen Bedingungen", 2)
    latency_rows = []
    for kind, title in [("chart", "30 Diagramme"), ("invoice", "20 Belege")]:
        lat = groups[kind + "_de_image"]["latency_seconds"]
        latency_rows.append([title, decimal(lat["median"]) + " s", decimal(lat["p95_nearest_rank"]) + " s"])
    table(doc, ["Deutscher Hauptlauf", "Median", "95. Perzentil"], latency_rows, [3.0, 1.75, 1.75])
    small(doc, "Nur Modellvorwärtslauf, Batch 1, CPU und sechs Threads; Laden, Kodierung und separate "
          "Aufwärmläufe sind ausgeschlossen. Das 95. Perzentil nutzt die Nearest-Rank-Methode. "
          "Keine allgemeine Durchsatz- oder Hardwarevergleichsaussage.")
    doc.add_heading("Prüfbare Reproduktion", 2)
    para(doc, "1. Quellen aus den festgelegten Revisionen mit scripts/acquire_sources.py beschaffen. "
         "2. Auswahl und Fragen mit scripts/build_benchmark.py erzeugen; Freeze und Bildhashes "
         "mit scripts/verify_freeze.py prüfen. 3. Den dokumentierten CPU-Lauf mit 786.432 Pixeln "
         "und visuellem Aufwärmen ausführen. 4. Unbearbeitete Ausgaben mit scripts/score_images.py "
         "bewerten und diesen Bericht mit report/build_report.py erstellen.")
    small(doc, "Die vollständigen Befehle und Abhängigkeiten stehen im Reproduktionspaket. "
          "scores.json enthält Feldantworten, Kontrollpaare, Konfusionszählungen und Scorer-Prüfwerte; "
          "predictions.metadata.json enthält Laufkonfiguration, Versionen und Abschlussstatus. "
          "report/build_manifest.json bindet den Bericht an diese Eingaben.")
    versions = m["packages"]
    small(doc, "Laufversionen: " + "; ".join(f"{k} {v}" for k, v in versions.items()) + ".")
    completed = datetime.fromisoformat(m["completed_at"]).astimezone(timezone.utc).strftime("%d.%m.%Y %H:%M:%S UTC")
    small(doc, f"Abschluss des Hauptlaufs: {completed}. Ein einzelner zusätzlicher Wiederholungstest "
          f"ergab {'dieselben' if m['repeatability_probe']['same_answers'] else 'andere'} Antworten; "
          f"maximale absolute Wahrscheinlichkeitsabweichung {m['repeatability_probe']['maximum_absolute_probability_delta']:.3g}. "
          "Dieser technische Einzeltest gehört nicht zu den 90 gewerteten Läufen und belegt keine allgemeine Reproduzierbarkeit.")
    doc.add_heading("Quellen der drei Diagrammartabweichungen", 2)
    cases = {c["case_id"]: c for c in read_jsonl(root / "benchmark/cases.jsonl")}
    chart_evidence = []
    for case_id in ["chart-001", "chart-002", "chart-003"]:
        c = cases[case_id]
        link(doc, f"{case_id} · Quell-ID {c['source_id']} · fixierte Diagrammquelle", chart_src["source_url"])
        small(doc, "SHA-256 " + c["image_sha256"], mono=True)
        chart_evidence.append({"case_id": case_id, "source_id": c["source_id"],
                               "source_revision": c["source_revision"], "image_sha256": c["image_sha256"]})

    # Page 6: documentary evidence expressly authorized for this report.
    evidence = make_evidence(root, destination.parent / "evidence")
    new_page(doc, "Bildbelege für die beiden Vorzeichenfehler")
    para(doc, "Die Summenblöcke zeigen die tatsächlich gedruckten negativen Gesamtbeträge zusammen mit "
         "Netto- und Umsatzsteuerbetrag. Die Ausschnitte sind unveränderte Bildpunkte der Testbilder: "
         "nur zugeschnitten, ohne Retusche oder neue Bildgenerierung. Die Darstellungsgröße im Bericht "
         "ist vergrößert. Die Referenzklasse bleibt in beiden Fällen „negativ“.")
    small(doc, "Quelle beider Abbildungen: Belege (Immineal, 2026), laterrr/belege-de-invoices-sample, "
          f"Revision {invoice_src['revision']}. Alle Belegdaten sind synthetisch.")
    for idx, e in enumerate(evidence, 1):
        doc.add_heading(f"Fall {e['case_id'].split('-')[1]} mit negativem Gesamtbetrag", 2)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = paragraph.add_run().add_picture(str(root / e["crop_file"]), width=Inches(5.5))
        picture._inline.docPr.set("descr", f"Originalausschnitt {e['source_id']}: Netto {e['amount_texts']['net_total']}, "
                                  f"Umsatzsteuer {e['amount_texts']['vat_amount_19']}, Gesamtbetrag {e['amount_texts']['gross_total']} Euro")
        link(doc, f"Abbildung {idx} · {e['case_id']} · Quell-ID {e['source_id']}", invoice_src["source_url"])
        f = next(f for f in gross_errors if f["case_id"] == e["case_id"])
        small(doc, f"Gedruckter Gesamtbetrag {e['amount_texts']['gross_total']} EUR. Modellantwort "
              f"„{LABELS[f['predicted']]}“ mit {percent(f['confidence'])} Konfidenz. Quelle: Belege (Immineal, 2026).")
        small(doc, "SHA-256 des vollständigen Testbilds\n" + e["source_image_sha256"], mono=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    manifest = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "output": destination.name,
        "output_sha256": sha(destination),
        "report_script_sha256": sha(__file__),
        "inputs_sha256": {
            str(scores_path.relative_to(root)): sha(scores_path),
            str(metadata_path.relative_to(root)): sha(metadata_path),
            "benchmark/freeze_manifest.json": sha(root / "benchmark/freeze_manifest.json"),
            "benchmark/design_summary.json": sha(root / "benchmark/design_summary.json"),
            "source/provenance.json": sha(root / "source/provenance.json"),
        },
        "intended_pages": 6,
        "visual_qa": "pending render and inspection of every page",
        "no_source_images_embedded": False,
        "documentary_figures_only": 2,
        "documentary_figure_provenance": evidence,
        "chart_ambiguity_provenance": chart_evidence,
    }
    (destination.parent / "build_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--scores", type=Path)
    parser.add_argument("--metadata", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--check-only", action="store_true", help="Validate final inputs without authoring")
    args = parser.parse_args()
    root = args.root.resolve()
    scores_path = (args.scores or root / "results/scores.json").resolve()
    metadata_path = (args.metadata or root / "results/predictions.metadata.json").resolve()
    output = args.output or root / "report/Clef_Bildbenchmark_2026-10-02.docx"
    try:
        scores, metadata, design, provenance = load_checked(root, scores_path, metadata_path)
    except (ValueError, FileNotFoundError, KeyError) as exc:
        print("REPORT NOT READY: " + str(exc), file=sys.stderr)
        return 2
    if args.check_only:
        print("Final report inputs validated; no document created")
        return 0
    out = create_report(scores, metadata, design, provenance, root, output, scores_path, metadata_path)
    print(f"Created {out}; mandatory next step: render and inspect every page")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
